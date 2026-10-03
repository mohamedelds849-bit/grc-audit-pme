# -*- coding: utf-8 -*-
"""
Build complet du livrable GRC « AgriSouss SARL » (étude de cas anonymisée).

    python build.py            # génère CSV/XLSX, graphiques, Markdown complet et PDF
    python build.py --no-pdf   # tout sauf le PDF

Entrées  : src/annexa.py (93 mesures), src/register.py (risques, actions),
           report/rapport_template.md (texte rédigé, avec marqueurs {{...}}).
Sorties  : outputs/*.csv, outputs/gap_analysis_annexA.xlsx, outputs/figures/*.png,
           report/rapport_audit_AgriSouss.md, outputs/rapport_audit_AgriSouss.pdf
Aucune API externe, aucun accès réseau.
"""
import csv
import datetime as dt
import os
import sys

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(ROOT, "src"))

import annexa  # noqa: E402
import register  # noqa: E402
import md2pdf  # noqa: E402

OUT = os.path.join(ROOT, "outputs")
FIG = os.path.join(OUT, "figures")
REPORT_DIR = os.path.join(ROOT, "report")
TEMPLATE = os.path.join(REPORT_DIR, "rapport_template.md")
MD_OUT = os.path.join(REPORT_DIR, "rapport_audit_AgriSouss.md")
PDF_OUT = os.path.join(OUT, "rapport_audit_AgriSouss.pdf")
REPORT_DATE = "octobre 2026"

# Palette (validated reference palette of the dataviz method + fixed status colours)
C_CUR, C_TGT = "#2a78d6", "#eb6834"
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e4e3df"
STATUS = {"Faible": "#0ca30c", "Modéré": "#fab219", "Élevé": "#ec835a", "Critique": "#d03b3b"}


def fmt_mad(n):
    return f"{n:,}".replace(",", " ")


# ------------------------------------------------------------------ tabular outputs
GAP_COLS = ["ID", "Thème", "Mesure (EN, ISO/IEC 27001:2022)", "Libellé FR (traduction de travail)",
            "Nouveau 2022", "Applicable", "État actuel", "Maturité (0-5)", "Cible", "Écart (points)",
            "Écart / recommandation", "Priorité", "Risques liés"]


def gap_rows():
    rows = []
    for r in annexa.controls():
        t = annexa.theme_of(r["id"])
        app = r["applicable"]
        rows.append([
            r["id"], annexa.THEMES_SHORT[t], r["title_en"], r["titre_fr"],
            "Oui" if r["id"] in annexa.NEW_2022 else "Non",
            "Oui" if app else "Non", r["etat"],
            r["maturite"] if app else "N/A", r["cible"] if app else "N/A",
            (r["cible"] - r["maturite"]) if app else "N/A",
            r["ecart"], r["priorite"], r["risques"],
        ])
    return rows


SOA_COLS = ["ID", "Mesure (EN)", "Applicable", "Justification inclusion / exclusion",
            "État de mise en oeuvre", "Action(s) du plan"]


def actions_for_control(cid):
    ids = [a["id"] for a in register.ACTIONS
           if cid in [m.strip() for m in a["mesures"].split(",")]]
    return ", ".join(ids)


def soa_rows():
    return [[r["id"], r["title_en"], "Oui" if r["applicable"] else "Non",
             annexa.soa_justification(r), annexa.soa_status(r), actions_for_control(r["id"])]
            for r in annexa.controls()]


RISK_COLS = ["ID", "Risque", "Source", "Cible (VM / bien support)", "Gravité", "Vraisemblance",
             "Score", "Niveau", "Mesures Annexe A", "Actions", "G résiduelle", "V résiduelle",
             "Score résiduel", "Niveau résiduel"]


def risk_rows():
    return [[r["id"], r["titre"], r["source"], r["cible"], r["g"], r["v"], r["score"], r["niveau"],
             r["mesures"], r["actions"], r["g_res"], r["v_res"], r["score_res"], r["niveau_res"]]
            for r in register.RISKS]


ACT_COLS = ["ID", "Horizon", "Action", "Effort", "Coût estimé (MAD HT)", "Porteur", "Risques", "Mesures"]


def action_rows():
    return [[a["id"], a["horizon"], a["titre"], a["effort"], a["cout"], a["porteur"], a["risques"], a["mesures"]]
            for a in register.ACTIONS]


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:   # BOM -> Excel lit les accents
        w = csv.writer(f, delimiter=";")
        w.writerow(header)
        w.writerows(rows)


def write_xlsx(path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import CellIsRule
    except ImportError:
        print("  [i] openpyxl absent : XLSX non généré (CSV disponibles)")
        return False
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(bold=True, color="FFFFFF")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Ecart_AnnexeA"
    sheet(ws, GAP_COLS, gap_rows(), [7, 14, 34, 32, 9, 10, 45, 10, 7, 9, 45, 8, 14])
    n = len(annexa.controls()) + 1
    for lo, hi, color in [(0, 0, "F4B6B6"), (1, 1, "F9D3A8"), (2, 2, "FCE9A6"), (3, 5, "BFE3BF")]:
        ws.conditional_formatting.add(f"H2:H{n}", CellIsRule(operator="between", formula=[str(lo), str(hi)],
                                                             fill=PatternFill("solid", fgColor=color)))
    sheet(wb.create_sheet("DdA_SoA"), SOA_COLS, soa_rows(), [7, 40, 10, 55, 18, 14])

    ws3 = wb.create_sheet("Synthese_themes")
    ts = annexa.theme_scores()
    rows = [[annexa.THEMES[t], v["n"], v["applicable"], v["actuel"], v["cible"]] for t, v in ts.items()]
    sheet(ws3, ["Thème", "Mesures", "Applicables", "Maturité actuelle", "Cible"], rows, [30, 10, 12, 16, 10])
    # formule globale pondérée (recalculée par Excel)
    last = len(rows) + 1
    ws3.append(["Global (pondéré)", f"=SUM(B2:B{last})", f"=SUM(C2:C{last})",
                f"=ROUND(SUMPRODUCT(C2:C{last},D2:D{last})/SUM(C2:C{last}),2)",
                f"=ROUND(SUMPRODUCT(C2:C{last},E2:E{last})/SUM(C2:C{last}),2)"])
    for c in ws3[last + 1]:
        c.font = Font(bold=True)

    sheet(wb.create_sheet("Risques"), RISK_COLS, risk_rows(), [6, 50, 26, 30, 8, 12, 7, 10, 18, 16, 10, 10, 10, 12])
    sheet(wb.create_sheet("Plan_action"), ACT_COLS, action_rows(), [6, 9, 70, 8, 22, 22, 14, 18])
    wb.save(path)
    return True


# ------------------------------------------------------------------ figures
def global_score():
    rows = [r for r in annexa.controls() if r["applicable"]]
    return round(sum(r["maturite"] for r in rows) / len(rows), 2)


def make_charts():
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        print("  [i] matplotlib absent : graphiques non générés")
        return {}
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "text.color": INK,
                         "axes.labelcolor": INK2, "xtick.color": INK2, "ytick.color": INK2,
                         "axes.edgecolor": GRID})
    ts = annexa.theme_scores()
    labels = [annexa.THEMES_SHORT[t] for t in ts][::-1]
    cur = [ts[t]["actuel"] for t in ts][::-1]
    tgt = [ts[t]["cible"] for t in ts][::-1]
    out = {}

    def bars(path, figsize, title):
        fig, ax = plt.subplots(figsize=figsize, dpi=200)
        y = range(len(labels))
        h = 0.36
        ax.barh([i + h / 2 + 0.02 for i in y], tgt, height=h, color=C_TGT, label="Cible à 12 mois")
        ax.barh([i - h / 2 - 0.02 for i in y], cur, height=h, color=C_CUR, label="Maturité actuelle")
        for i in y:
            ax.text(cur[i] + 0.05, i - h / 2 - 0.02, f"{cur[i]:.2f}".replace(".", ","), va="center", fontsize=8.5, color=INK)
            ax.text(tgt[i] + 0.05, i + h / 2 + 0.02, f"{tgt[i]:.2f}".replace(".", ","), va="center", fontsize=8.5, color=INK2)
        ax.set_yticks(list(y))
        ax.set_yticklabels(labels)
        ax.set_xlim(0, 5)
        ax.set_xticks(range(6))
        ax.set_xlabel("Maturité moyenne des mesures applicables (0 = inexistant, 5 = optimisé)")
        ax.grid(axis="x", color=GRID, linewidth=0.8)
        ax.set_axisbelow(True)
        for s in ("top", "right", "left"):
            ax.spines[s].set_visible(False)
        ax.tick_params(axis="y", length=0)
        handles, lbls = ax.get_legend_handles_labels()
        ax.legend(handles[::-1], lbls[::-1], frameon=False, loc="lower right", fontsize=8.5)
        ax.set_title(title, loc="left", fontsize=11, color=INK, pad=10)
        fig.tight_layout()
        fig.savefig(path)
        plt.close(fig)

    p = os.path.join(FIG, "maturite_par_theme.png")
    bars(p, (7.2, 3.4), "Maturité ISO/IEC 27001:2022 Annexe A par thème - AgriSouss SARL")
    out["maturity"] = p
    p = os.path.join(FIG, "maturite_synthese.png")
    bars(p, (6.4, 2.5), "Maturité par thème : actuelle vs cible")
    out["small"] = p

    # matrice des risques : cellules colorées par niveau + libellé du niveau (jamais la couleur seule)
    fig, ax = plt.subplots(figsize=(6.6, 4.8), dpi=200)
    for v in range(1, 5):
        for g in range(1, 5):
            _, lvl = register.niveau(g, v)
            ax.add_patch(plt.Rectangle((g - 0.5 + 0.02, v - 0.5 + 0.02), 0.96, 0.96,
                                       facecolor=STATUS[lvl], alpha=0.28, edgecolor="white", linewidth=2))
            ids = [r["id"] for r in register.RISKS if r["g"] == g and r["v"] == v]
            ax.text(g - 0.44, v + 0.40, lvl, fontsize=6.5, color=INK2, va="top")
            if ids:
                txt = "\n".join(", ".join(ids[i:i + 2]) for i in range(0, len(ids), 2))
                ax.text(g, v - 0.05, txt, ha="center", va="center", fontsize=9, color=INK, fontweight="bold")
    ax.set_xlim(0.5, 4.5)
    ax.set_ylim(0.5, 4.5)
    ax.set_xticks(range(1, 5))
    ax.set_xticklabels([f"{k}\n{v}" for k, v in register.GRAVITE.items()], fontsize=8)
    ax.set_yticks(range(1, 5))
    ax.set_yticklabels([f"{k} {v}" for k, v in register.VRAISEMBLANCE.items()], fontsize=8)
    ax.set_xlabel("Gravité")
    ax.set_ylabel("Vraisemblance")
    for s in ax.spines.values():
        s.set_visible(False)
    ax.tick_params(length=0)
    ax.set_title("Matrice des risques avant traitement (12 risques)", loc="left", fontsize=11, pad=10)
    fig.tight_layout()
    p = os.path.join(FIG, "matrice_risques.png")
    fig.savefig(p)
    plt.close(fig)
    out["risk"] = p
    return out


# ------------------------------------------------------------------ markdown blocks
def md_table(header, rows):
    esc = lambda x: str(x).replace("|", "/").replace("\n", " ")
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    lines += ["| " + " | ".join(esc(c) for c in row) + " |" for row in rows]
    return "\n".join(lines)


def build_markdown(charts):
    ts = annexa.theme_scores()
    ctrls = annexa.controls()
    budget = register.budget_by_horizon()
    rel = lambda p: os.path.relpath(p, REPORT_DIR).replace("\\", "/")

    def img(key, caption, width):
        if key not in charts:
            return f"*(Graphique « {caption} » non généré : matplotlib indisponible.)*"
        return f"![{caption}]({rel(charts[key])}){{width={width}}}"

    rep = {}
    rep["DATE"] = REPORT_DATE
    rep["SCORE_GLOBAL"] = f"{global_score():.1f}".replace(".", ",")
    app = [r for r in annexa.controls() if r["applicable"]]
    rep["SCORE_TARGET"] = f"{sum(r['cible'] for r in app) / len(app):.1f}".replace(".", ",")
    rep["NB_RISK_CRIT"] = str(sum(r["niveau"] == "Critique" for r in register.RISKS))
    rep["NB_RISK_HIGH"] = str(sum(r["niveau"] == "Élevé" for r in register.RISKS))
    rng = lambda h: f"{fmt_mad(budget[h][0])} à {fmt_mad(budget[h][1])}"
    rep["BUDGET_30"], rep["BUDGET_90"], rep["BUDGET_12"] = rng("30 j"), rng("90 j"), rng("12 mois")
    rep["CHART_SMALL"] = img("small", "Maturité par thème : actuelle vs cible", "13cm")
    rep["CHART_MATURITY"] = img("maturity", "Figure 1 - Maturité Annexe A par thème", "15.5cm")
    rep["CHART_RISK"] = img("risk", "Figure 2 - Matrice des risques avant traitement", "12.5cm")

    rep["RISK_REGISTER"] = md_table(
        ["Réf.", "Risque", "Source de risque", "G", "V", "Niveau", "Mesures Annexe A"],
        [[r["id"], r["titre"], r["source"], r["g"], r["v"], f"{r['score']} {r['niveau']}", r["mesures"]]
         for r in register.RISKS])

    rows = []
    for v in range(4, 0, -1):
        row = [f"V{v} {register.VRAISEMBLANCE[v]}"]
        for g in range(1, 5):
            ids = [r["id"] for r in register.RISKS if r["g"] == g and r["v"] == v]
            row.append(" ".join(ids) if ids else "-")
        rows.append(row)
    rep["RISK_MATRIX_TABLE"] = md_table(["Vraisemblance \\ Gravité"] + [f"G{g} {register.GRAVITE[g]}" for g in range(1, 5)], rows)

    rep["RISK_RESIDUAL_TABLE"] = md_table(
        ["Réf.", "Niveau initial", "Actions", "Niveau résiduel visé", "Décision"],
        [[r["id"], f"{r['score']} {r['niveau']}", r["actions"], f"{r['score_res']} {r['niveau_res']}",
          "Accepté après traitement" if r["score_res"] <= 3 else "Réduit puis accepté (surveillance)"]
         for r in register.RISKS])

    tot_app = sum(v["applicable"] for v in ts.values())
    rows = [[annexa.THEMES[t], v["n"], v["applicable"], f"{v['actuel']:.2f}", f"{v['cible']:.2f}",
             f"{v['cible'] - v['actuel']:.2f}"] for t, v in ts.items()]
    g_cur = sum(v["actuel"] * v["applicable"] for v in ts.values()) / tot_app
    g_tgt = sum(v["cible"] * v["applicable"] for v in ts.values()) / tot_app
    rows.append(["**Global (pondéré)**", "**93**", f"**{tot_app}**", f"**{g_cur:.2f}**", f"**{g_tgt:.2f}**",
                 f"**{g_tgt - g_cur:.2f}**"])
    rep["THEME_TABLE"] = md_table(["Thème", "Mesures", "Applicables", "Maturité actuelle", "Cible 12 mois", "Écart"], rows)

    p1 = [r for r in ctrls if r["priorite"] == "P1"]
    rep["TOP_GAPS"] = (f"{len(p1)} mesures sont classées P1 (à traiter dans les 90 jours) :\n\n" + md_table(
        ["Mesure", "Constat", "Recommandation", "Mat.", "Risques"],
        [[f"{r['id']} {r['titre_fr']}", r["etat"], r["ecart"], f"{r['maturite']}→{r['cible']}", r["risques"]] for r in p1]))

    parts = []
    names = {"30 j": "Horizon 1 - Quick wins (0 à 30 jours)", "90 j": "Horizon 2 - Structurer (30 à 90 jours)",
             "12 mois": "Horizon 3 - Pérenniser (3 à 12 mois)"}
    for h, title in names.items():
        acts = [a for a in register.ACTIONS if a["horizon"] == h]
        parts.append(f"## {title}\n\n" + md_table(
            ["Réf.", "Action", "Effort", "Coût estimé (MAD HT)", "Porteur", "Risques"],
            [[a["id"], a["titre"], a["effort"], a["cout"], a["porteur"], a["risques"]] for a in acts]))
    rep["ACTION_PLAN"] = "\n\n".join(parts)

    lo = sum(b[0] for b in budget.values())
    hi = sum(b[1] for b in budget.values())
    rep["BUDGET_TABLE"] = md_table(["Horizon", "Fourchette basse (MAD)", "Fourchette haute (MAD)"],
                                   [[names[h].split(" - ")[1], fmt_mad(budget[h][0]), fmt_mad(budget[h][1])] for h in names]
                                   + [["**Total première année**", f"**{fmt_mad(lo)}**", f"**{fmt_mad(hi)}**"]]) + \
        "\n\nLes postes « / an » (gestionnaire de mots de passe, messagerie professionnelle, CDN) sont récurrents."

    rep["NB_APPLICABLE"] = str(sum(r["applicable"] for r in ctrls))
    rep["NB_EXCLUDED"] = str(sum(not r["applicable"] for r in ctrls))
    soa_parts = []
    for t, name in annexa.THEMES.items():
        rows = [[r[0], r[1], r[2], r[3], r[4], r[5] or "-"] for r in soa_rows() if r[0].split(".")[0] == t]
        soa_parts.append(f"## A.{t} {name}\n\n" + md_table(["ID", "Mesure (EN)", "Appl.", "Justification", "État", "Actions"], rows))
    rep["SOA"] = "\n\n".join(soa_parts)

    with open(TEMPLATE, encoding="utf-8") as f:
        text = f.read()
    for k, v in rep.items():
        text = text.replace("{{" + k + "}}", v)
    left = [w for w in text.split("{{")[1:]]
    if left:
        raise SystemExit("Marqueurs non remplacés : " + ", ".join(x.split("}}")[0] for x in left))
    header = ("<!-- FICHIER GÉNÉRÉ par build.py à partir de report/rapport_template.md, src/annexa.py et "
              "src/register.py. Modifier les sources, pas ce fichier. -->\n")
    # le front-matter doit rester en tête : on insère le commentaire après lui
    fm_end = text.index("---", 3) + 3
    text = text[:fm_end] + "\n" + header + text[fm_end:]
    with open(MD_OUT, "w", encoding="utf-8") as f:
        f.write(text)
    return text


# ------------------------------------------------------------------ main
def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    os.makedirs(FIG, exist_ok=True)
    print("[1/5] Tableaux CSV")
    write_csv(os.path.join(OUT, "gap_analysis_annexA.csv"), GAP_COLS, gap_rows())
    write_csv(os.path.join(OUT, "soa_draft.csv"), SOA_COLS, soa_rows())
    write_csv(os.path.join(OUT, "risk_register.csv"), RISK_COLS, risk_rows())
    write_csv(os.path.join(OUT, "action_plan.csv"), ACT_COLS, action_rows())
    print("[2/5] Classeur XLSX :", "ok" if write_xlsx(os.path.join(OUT, "gap_analysis_annexA.xlsx")) else "ignoré")
    print("[3/5] Graphiques")
    charts = make_charts()
    print("[4/5] Markdown complet ->", os.path.relpath(MD_OUT, ROOT))
    build_markdown(charts)
    if "--no-pdf" in sys.argv:
        return
    print("[5/5] PDF ->", os.path.relpath(PDF_OUT, ROOT))
    md2pdf.render(MD_OUT, PDF_OUT)
    try:
        from pypdf import PdfReader
        n = len(PdfReader(PDF_OUT).pages)
        print(f"      vérification pypdf : {n} pages")
    except ImportError:
        print("      (pypdf absent : vérification du nombre de pages ignorée)")
    print(f"Score global de maturité : {global_score():.2f} / 5")


if __name__ == "__main__":
    main()
